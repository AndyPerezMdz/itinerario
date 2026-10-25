import openpyxl, sys, re, collections, datetime as dt
MESES=["ENERO","FEBRERO","MARZO","ABRIL","MAYO","JUNIO","JULIO","AGOSTO","SEPTIEMBRE","OCTUBRE","NOVIEMBRE","DICIEMBRE"]
def extraer(path):
    wb=openpyxl.load_workbook(path,data_only=True)
    out=[]
    for ws in wb.worksheets:
        mes=MESES.index(ws.title.strip().upper())+1
        r=4
        while r<=ws.max_row:
            fechas=[ws.cell(r,c).value for c in range(2,9)]
            if not all(isinstance(f,dt.datetime) for f in fechas): r+=1; continue
            for rr in range(r+1,r+5):
                for i,c in enumerate(range(2,9)):
                    v=ws.cell(rr,c).value
                    f=fechas[i].date()
                    if isinstance(v,str) and v.strip() and f.month==mes:
                        out.append((f,re.sub(r"\s+"," ",v.strip().upper())))
            r+=5
    return sorted(set(out))
if __name__=="__main__":
    ev=extraer(sys.argv[1])
    print(len(ev), ev[0][0], ev[-1][0])
    cnt=collections.Counter(t for _,t in ev)
    dias="LMXJVSD"
    for t,n in cnt.most_common():
        wd=collections.Counter(dias[f.weekday()] for f,tt in ev if tt==t)
        print(n, t, dict(wd))
